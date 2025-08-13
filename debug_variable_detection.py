#!/usr/bin/env python3
"""
Debug Variable Detection
Test to see what variables are actually being detected vs. what should be detected.
"""

import re
from docx import Document
from document_parser import DocumentParser
from variable_utils import extract_all_variables
from logging_config import logger

def debug_variable_detection():
    """Debug variable detection to see what's being missed."""
    
    print("🔍 DEBUGGING VARIABLE DETECTION")
    print("="*70)
    
    # Test the actual model document from the logs
    model_path = "C:/Users/brand/Downloads/project/ModUpate/assets/templates/DRAFT_CY2026_8_PPO_MA_EOC_FINAL.docx"
    
    try:
        # Method 1: Direct DOCX parsing (what smart_wizard uses)
        print("\n📄 METHOD 1: Direct DOCX parsing")
        doc = Document(model_path)
        all_variables_method1 = []
        provider_related_paras = []
        
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue
                
            # Check for provider-related content
            if "provider" in text.lower() or "www." in text.lower() or "website" in text.lower():
                provider_related_paras.append({
                    "para_index": i,
                    "text": text,
                    "variables": list(re.findall(r'\[([^\]]+)\]', text))
                })
                
            # Find all variables in this paragraph
            variables = list(re.finditer(r'\[([^\]]+)\]', text))
            for match in variables:
                var_name = match.group(1).strip()
                all_variables_method1.append({
                    "name": var_name,
                    "paragraph": i,
                    "context": text[:50] + "..." if len(text) > 50 else text
                })
        
        print(f"   Found {len(all_variables_method1)} variables total")
        
        # Method 2: Document parser (what main system uses)
        print("\n📋 METHOD 2: DocumentParser class")
        parser = DocumentParser()
        parsed_blocks = parser.parse(model_path)
        variable_blocks = [block for block in parsed_blocks if block.get('type') == 'variable']
        
        print(f"   Parsed {len(parsed_blocks)} total blocks")
        print(f"   Found {len(variable_blocks)} variable blocks")
        
        # Method 3: Check provider directory specifically
        print("\n🌐 PROVIDER DIRECTORY ANALYSIS")
        print(f"Found {len(provider_related_paras)} provider-related paragraphs:")
        
        for para_info in provider_related_paras[:5]:  # Show first 5
            print(f"   Para {para_info['para_index']}: '{para_info['text'][:100]}...'")
            if para_info['variables']:
                print(f"      Variables: {para_info['variables']}")
            else:
                print(f"      No variables detected")
                # Check if there are URLs that should be variables
                if "www." in para_info['text'] or "http" in para_info['text']:
                    print(f"      ⚠️  Contains URL but no [insert URL] variable!")
        
        # Method 4: Search for specific patterns that might be missed
        print("\n🔍 MISSING PATTERN ANALYSIS")
        missed_patterns = []
        
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue
                
            # Look for URL patterns that might not have proper variables
            if ("www." in text or "http" in text) and "[insert url]" not in text.lower():
                missed_patterns.append({
                    "type": "URL without variable",
                    "paragraph": i,
                    "text": text
                })
            
            # Look for patterns that look like they should be variables but aren't bracketed
            potential_vars = [
                r"insert\s+[a-zA-Z]+",
                r"add\s+[a-zA-Z]+", 
                r"include\s+[a-zA-Z]+",
                r"www\.[a-zA-Z0-9.-]+",
                r"https?://[a-zA-Z0-9.-]+"
            ]
            
            for pattern in potential_vars:
                matches = re.findall(pattern, text, re.IGNORECASE)
                if matches and not re.search(r'\[.*\]', text):
                    missed_patterns.append({
                        "type": f"Unbracketed pattern: {pattern}",
                        "paragraph": i,
                        "matches": matches,
                        "text": text[:100] + "..." if len(text) > 100 else text
                    })
        
        print(f"Found {len(missed_patterns)} potential missed patterns:")
        for pattern in missed_patterns[:10]:  # Show first 10
            print(f"   {pattern['type']}: Para {pattern['paragraph']}")
            print(f"      Text: {pattern['text']}")
            if 'matches' in pattern:
                print(f"      Matches: {pattern['matches']}")
        
        # Summary comparison
        print("\n📊 SUMMARY")
        print(f"Method 1 (Direct): {len(all_variables_method1)} variables")
        print(f"Method 2 (Parser): {len(variable_blocks)} variables")
        print(f"Provider-related paragraphs: {len(provider_related_paras)}")
        print(f"Potential missed patterns: {len(missed_patterns)}")
        
        # Check if specific problematic variables exist
        problematic_vars = [
            "insert URL",
            "www.bcbsm.com/providersmedicare",
            "provider directory",
            "website"
        ]
        
        print("\n🎯 SPECIFIC VARIABLE CHECK")
        found_vars = [var['name'].lower() for var in all_variables_method1]
        
        for prob_var in problematic_vars:
            if any(prob_var.lower() in fvar for fvar in found_vars):
                print(f"   ✅ Found: {prob_var}")
            else:
                print(f"   ❌ Missing: {prob_var}")
        
        return len(all_variables_method1), len(variable_blocks), len(missed_patterns)
        
    except Exception as e:
        print(f"❌ Error during debugging: {e}")
        import traceback
        traceback.print_exc()
        return 0, 0, 0

def enhanced_variable_detection_test():
    """Test enhanced variable detection patterns."""
    
    print("\n🔧 ENHANCED VARIABLE DETECTION TEST")
    print("="*70)
    
    # Test more sophisticated patterns
    enhanced_patterns = [
        r'\[([^\]]+)\]',  # Standard brackets
        r'\{([^}]+)\}',   # Curly brackets (sometimes used)
        r'<<([^>]+)>>',   # Double angle brackets
        r'\(\(([^)]+)\)\)',  # Double parentheses
        r'INSERT\s*:\s*([^.\n]+)',  # INSERT: pattern
        r'ADD\s*:\s*([^.\n]+)',     # ADD: pattern
        r'INCLUDE\s*:\s*([^.\n]+)', # INCLUDE: pattern
    ]
    
    test_text = """
    The most recent list of providers is also available on our website at www.bcbsm.com/providersmedicare.
    [Insert as applicable:] We included a copy of our Provider Directory in the envelope with this document.
    INSERT: Plan specific information here.
    ADD: Additional content as needed.
    {variable in curly brackets}
    <<variable in angle brackets>>
    ((variable in double parens))
    """
    
    print("Testing enhanced patterns on sample text:")
    print(f"Sample: {test_text[:100]}...")
    
    for i, pattern in enumerate(enhanced_patterns):
        matches = re.findall(pattern, test_text, re.IGNORECASE)
        print(f"   Pattern {i+1} ({pattern}): {len(matches)} matches")
        if matches:
            print(f"      Found: {matches}")
    
    return len(enhanced_patterns)

if __name__ == "__main__":
    try:
        vars_method1, vars_method2, missed_patterns = debug_variable_detection()
        enhanced_patterns = enhanced_variable_detection_test()
        
        print(f"\n{'='*70}")
        print("🎯 DEBUGGING CONCLUSIONS:")
        
        if vars_method1 > 0:
            print(f"✅ Variable detection is working (found {vars_method1} variables)")
            if missed_patterns > 10:
                print(f"⚠️  However, {missed_patterns} potential patterns might be missed")
                print("💡 Consider adding enhanced pattern detection")
            else:
                print("✅ Pattern detection seems comprehensive")
        else:
            print("❌ Variable detection has major issues")
            
        if abs(vars_method1 - vars_method2) > 10:
            print(f"⚠️  Discrepancy between detection methods: {vars_method1} vs {vars_method2}")
        else:
            print("✅ Detection methods are consistent")
            
        print("\n🚀 RECOMMENDATIONS:")
        print("1. Check if document path is correct")
        print("2. Verify variable format in source document")
        print("3. Consider enhanced pattern detection for edge cases")
        print("4. Test with actual document from logs")
        
    except Exception as e:
        print(f"❌ Debug script failed: {e}")
        import traceback
        traceback.print_exc()