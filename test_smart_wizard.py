#!/usr/bin/env python3
"""
Test script for the Smart Wizard with mock data.
"""

import tkinter as tk
from smart_wizard import SmartWizard
import tempfile
import os
from docx import Document

def create_test_docx():
    """Create a test DOCX file with variables."""
    doc = Document()
    
    # Add paragraphs with variables
    doc.add_paragraph("This plan, [insert 2026 plan name], is offered by [insert MAO name].")
    doc.add_paragraph("Call [insert phone number] or TTY [insert TTY number] for help.")
    doc.add_paragraph("Coverage is available in [insert state] and surrounding areas.")
    doc.add_paragraph("Premium information: [insert monthly premium amount].")
    doc.add_paragraph("Contact us at [insert customer service number] for questions.")
    
    # Save to temp file
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
    doc.save(temp_file.name)
    return temp_file.name

def test_smart_wizard():
    """Test the smart wizard with mock data."""
    
    # Create mock extracted values
    extracted_values = {
        "insert 2026 plan name": "Medicare Plus Blue Group PPO",
        "insert MAO name": "Blue Cross Blue Shield of Michigan", 
        "insert phone number": "1-855-669-8040",
        "insert TTY number": "711",
        "insert state": "Michigan",
        "insert monthly premium amount": "None",
        "insert customer service number": "1-800-555-0199"
    }
    
    # Create test DOCX
    test_docx_path = create_test_docx()
    
    try:
        # Create root window
        root = tk.Tk()
        root.title("Test Dashboard")
        root.geometry("400x300")
        
        # Add a test button to launch wizard
        def launch_wizard():
            try:
                wizard = SmartWizard(
                    parent=root,
                    extracted_values=extracted_values,
                    model_docx_path=test_docx_path
                )
                print(f"✅ Smart Wizard launched successfully!")
                print(f"📊 Variables found: {len(wizard.variables)}")
                for var in wizard.variables:
                    print(f"   - {var['name']} ({var['type']})")
                    
            except Exception as e:
                print(f"❌ Error launching wizard: {e}")
                import traceback
                traceback.print_exc()
        
        test_button = tk.Button(
            root,
            text="🧙‍♂️ Launch Smart Wizard",
            command=launch_wizard,
            font=('Arial', 14),
            bg='#4a90e2',
            fg='white',
            padx=20,
            pady=10
        )
        test_button.pack(expand=True)
        
        # Add info label
        info_label = tk.Label(
            root,
            text=f"Test Data:\n{len(extracted_values)} extracted values\nSample variables ready",
            font=('Arial', 10),
            bg='white',
            fg='#333'
        )
        info_label.pack(pady=20)
        
        print("🧪 Smart Wizard Test")
        print("=" * 40)
        print(f"📄 Test DOCX created: {test_docx_path}")
        print(f"📊 Mock extracted values: {len(extracted_values)}")
        print("🖱️  Click the button to test the wizard")
        print()
        
        # Start the GUI
        root.mainloop()
        
    finally:
        # Clean up temp file
        try:
            os.unlink(test_docx_path)
        except:
            pass

if __name__ == "__main__":
    test_smart_wizard() 