#!/usr/bin/env python3
"""
MBC Scaffolder

Builds or augments a context pack with precise MBC row locations and, optionally,
writes human-visible placeholders into a copy of the 2026 model so SMEs can see
where values will be filled. The filler can then rely on row coordinates (loc)
instead of bracket text to guarantee fill.

Usage (programmatic):
    from scaffold_mbc import scaffold_mbc
    result = scaffold_mbc(model_docx, base_pack_json, out_pack_json, scaffold_out_docx)

CLI:
    python scaffold_mbc.py <2026_model.docx> <in_pack.json> <out_pack.json> [scaffolded_out.docx]
"""

from __future__ import annotations

import json
import os
import re
from typing import Dict, List, Tuple

from docx import Document


def _slugify(text: str) -> str:
    t = re.sub(r"[^a-zA-Z0-9]+", "_", (text or "").lower())
    t = re.sub(r"_+", "_", t).strip("_")
    return t[:120]


def _is_generic_label(label: str) -> bool:
    words = (label or '').strip().split()
    low = (label or '').lower()
    if len(words) <= 2:
        return True
    generics = ['insert', 'list', 'as appropriate', 'as applicable', 'cost share']
    return any(g in low for g in generics)


def _effective_label(current: str, last_meaningful: str) -> Tuple[str, bool]:
    if _is_generic_label(current):
        return (last_meaningful or current or ''), True
    return (current or '', False)


def _load_pack(path: str) -> Dict:
    if not (path and os.path.exists(path)):
        return {"items": []}
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def _index_pack_items_by_canonical(pack: Dict) -> Dict[str, Dict]:
    out: Dict[str, Dict] = {}
    for it in pack.get('items', []):
        key = it.get('canonical_key') or it.get('name')
        if key:
            out.setdefault(key, it)
    return out


def scaffold_mbc(model_docx_path: str, in_pack_json_path: str, out_pack_json_path: str,
                 scaffolded_docx_out: str | None = None) -> Dict:
    doc = Document(model_docx_path)
    pack = _load_pack(in_pack_json_path)
    items = pack.get('items', []) if isinstance(pack, dict) else []
    by_key = _index_pack_items_by_canonical(pack)

    def add_occurrence(it: Dict, placeholder: str | None, table_idx: int, row_idx: int) -> None:
        occs = it.setdefault('occurrences', []) or []
        occs.append({
            'placeholder': placeholder,
            'loc': { 'table_idx': table_idx, 'row_idx': row_idx, 'cell_idx': 1 },
            'virtual': placeholder is None
        })

    # Optionally prepare scaffolded copy
    scaffold_doc = None
    if scaffolded_docx_out:
        scaffold_doc = Document(model_docx_path)

    created = 0
    updated = 0

    table_index = 0
    for tbl in doc.tables:
        try:
            if not tbl.rows or max(len(r.cells) for r in tbl.rows) < 2:
                table_index += 1
                continue
            header_text = ' '.join(c.text or '' for c in tbl.rows[0].cells)
            low = header_text.lower()
            looks_mbc = ('what you pay' in low) and any(k in low for k in ['service', 'services', 'benefit', 'benefits', 'coverage'])
            if not looks_mbc:
                table_index += 1
                continue
            last_meaningful_label = ''
            for ri, r in enumerate(tbl.rows[1:], start=1):
                cells = r.cells
                if len(cells) < 2:
                    continue
                left = (cells[0].text or '').strip()
                right = (cells[1].text or '').strip()
                eff_label, inherited = _effective_label(left, last_meaningful_label)
                if not _is_generic_label(left):
                    last_meaningful_label = left
                if not eff_label:
                    continue
                canonical = f"mbc_{_slugify(eff_label)}_cost_share"
                # Detect placeholder in right cell (first bracket)
                ph = None
                m = re.search(r"\[([^\]]+)\]", right)
                if m:
                    ph = m.group(0)
                # Ensure pack item exists
                it = by_key.get(canonical)
                if it is None:
                    # Create a minimal item
                    ctx = f"Service: {left} | What you pay: {'📍 ' + ph if ph else '—'}".strip()
                    it = {
                        'name': (m.group(1).strip() if m else f"MBC – {eff_label} cost share"),
                        'type': 'financial',
                        'vpx_kind': 'mbc_cost_share',
                        'canonical_key': canonical,
                        'full_context': ctx,
                        'before_context': left,
                        'after_context': '',
                        'in_table': True,
                        'occurrences': [],
                        'para_index': 0,
                        'mbc_effective_label': eff_label,
                        'mbc_generic_left': inherited,
                        'mbc_table_header': header_text,
                        'table_index': table_index,
                        'row_index': ri,
                    }
                    items.append(it)
                    by_key[canonical] = it
                    created += 1
                else:
                    updated += 1
                # Append occurrence with coordinates
                add_occurrence(it, ph, table_index, ri)
                # If scaffolding requested and no bracket present, write a light placeholder in copy
                if scaffold_doc is not None and ph is None:
                    try:
                        s_tbl = scaffold_doc.tables[table_index]
                        s_cell = s_tbl.rows[ri].cells[1]
                        placeholder = f"[[MBC – {eff_label} cost share]]"
                        s_cell.text = placeholder
                    except Exception:
                        pass
            table_index += 1
        except Exception:
            table_index += 1
            continue

    # Save scaffolded copy if requested
    if scaffold_doc is not None and scaffolded_docx_out:
        scaffold_doc.save(scaffolded_docx_out)

    # Save augmented pack
    out_pack = dict(pack)
    out_pack['items'] = items
    os.makedirs(os.path.dirname(out_pack_json_path) or '.', exist_ok=True)
    with open(out_pack_json_path, 'w', encoding='utf-8') as f:
        json.dump(out_pack, f, indent=2)

    return {
        'created_items': created,
        'updated_items': updated,
        'out_pack_path': out_pack_json_path,
        'scaffolded_docx_out': scaffolded_docx_out
    }


if __name__ == '__main__':
    import sys
    if len(sys.argv) < 4:
        print('Usage: python scaffold_mbc.py <2026_model.docx> <in_pack.json> <out_pack.json> [scaffolded_out.docx]')
        sys.exit(2)
    model = sys.argv[1]
    in_pack = sys.argv[2]
    out_pack = sys.argv[3]
    out_doc = sys.argv[4] if len(sys.argv) >= 5 else None
    res = scaffold_mbc(model, in_pack, out_pack, out_doc)
    print('[OK]', json.dumps(res, indent=2))


