import sys
import re
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("ERROR: python-docx not installed. Run: pip install python-docx")
    sys.exit(1)


def iter_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    yield p


def main():
    if len(sys.argv) < 2:
        print("Usage: python scan_docx.py <path-to-docx>")
        sys.exit(2)

    path = Path(sys.argv[1])
    print(f"PATH: {path}")
    print(f"EXISTS: {path.exists()}")
    if not path.exists():
        sys.exit(3)

    doc = Document(str(path))

    placeholders = []
    insert_like = []
    for para in iter_paragraphs(doc):
        txt = para.text or ""
        phs = re.findall(r"\[[^\]]+\]", txt)
        placeholders.extend(phs)
        for ph in phs:
            if re.search(r"\binsert\b", ph, flags=re.IGNORECASE):
                insert_like.append(ph)

    from collections import Counter
    cnt = Counter([s.strip() for s in placeholders])
    top = sorted(cnt.items(), key=lambda x: -x[1])[:30]
    print(f"TOTAL_PLACEHOLDERS: {len(placeholders)}")
    print(f"UNIQUE_PLACEHOLDERS: {len(cnt)}")
    print("TOP_PLACEHOLDERS:")
    for s, n in top:
        print(f"  {n} x {s}")
    print(f"INSERT_LIKE_COUNT: {len(insert_like)}")
    print("SAMPLE_INSERTS:")
    for s in list(dict.fromkeys(insert_like))[:30]:
        print(f"  {s}")


if __name__ == "__main__":
    main()


